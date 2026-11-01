from django.core.management.base import BaseCommand
from openpyxl import load_workbook
from datetime import datetime
from decimal import Decimal
from djangoapp.models import CartList, OilFee
import os


class Command(BaseCommand):
    help = "导入 Excel 数据到数据库"

    def add_arguments(self, parser):
        # 接收命令行参数
        parser.add_argument("file_type", type=str, help="要导入的数据类型: cart 或 oil")
        parser.add_argument("excel_path", type=str, help="Excel 文件的绝对路径")

    def handle(self, *args, **options):
        file_type = options["file_type"]
        excel_path = options["excel_path"]

        if not os.path.exists(excel_path):
            self.stdout.write(self.style.ERROR(f"找不到文件: {excel_path}"))
            return

        wb = load_workbook(excel_path, data_only=True)
        ws = wb.active

        if file_type == "cart":
            self.import_cart(ws)
        elif file_type == "oil":
            self.import_oil(ws)
        else:
            self.stdout.write(self.style.ERROR("不支持的类型，请用 'cart' 或 'oil'"))

    def import_cart(self, ws):
        """导入车辆"""
        created_count, updated_count = 0, 0
        # min_row=2 跳过表头，values_only=True 直接拿单元格的值
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]: continue  # 如果第一列(车牌)为空，跳过

            plate, admin, category, brand, p_date, service, tax, fee = row

            # 处理日期格式
            if isinstance(p_date, datetime):
                p_date = p_date.date()

            obj, created = CartList.objects.update_or_create(
                plate_number=str(plate).strip(),
                defaults={
                    "admin": admin,
                    "category": category,
                    "brand_model": brand,
                    "purchase_date": p_date,
                    "service_life": service or 15,
                    "use_tax": Decimal(str(tax or 0)),
                    "check_fee": Decimal(str(fee or 0))
                }
            )
            if created:
                created_count += 1
            else:
                updated_count += 1

        self.stdout.write(self.style.SUCCESS(f"✅ 车辆导入完成: 新增 {created_count} 条, 更新 {updated_count} 条"))

    def import_oil(self, ws):
        """导入油费"""
        created_count, skip_count = 0, 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]: continue

            plate, year, month, number, price, amount = row

            # 找外键车辆
            try:
                cart = CartList.objects.get(plate_number=str(plate).strip())
            except CartList.DoesNotExist:
                self.stdout.write(self.style.WARNING(f"⚠️ 跳过: 车辆 {plate} 不存在"))
                skip_count += 1
                continue

            # 使用 update_or_create 防止同车同月重复导入报错
            OilFee.objects.update_or_create(
                cart=cart, year=year, month=month,
                defaults={
                    "number": Decimal(str(number or 0)),
                    "price": Decimal(str(price or 0)),
                    "amount": Decimal(str(amount or 0))
                }
            )
            created_count += 1

        self.stdout.write(self.style.SUCCESS(f"✅ 油费导入完成: 处理 {created_count} 条, 跳过 {skip_count} 条"))